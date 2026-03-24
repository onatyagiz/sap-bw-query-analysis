"""
SAP BW Query Analysis — Master Script

Prompts for two SAP BW query Excel exports, runs all analyses in sequence,
and produces a single combined Excel workbook containing:

  Sheet 1  Query A (Filled)      — hierarchy columns A–N filled downward
  Sheet 2  Query B (Filled)      — hierarchy columns A–N filled downward
  Sheet 3  Comparison            — element-by-element diff (colour-coded)
  Sheet 4  Comparison Summary    — element counts per section
  Sheet 5  Restricted KFs        — RKF structural comparison
  Sheet 6  Calculated KFs        — CKF formula comparison
  Sheet 7  Formulas              — Row_XXXXX formula comparison
  Sheet 8  KF Summary            — key figure counts
"""

import os
import sys
import tempfile
from copy import copy
from pathlib import Path

import openpyxl

# ── Sibling modules ──────────────────────────────────────────────────────────

SCRIPT_DIR = Path(__file__).parent
sys.path.insert(0, str(SCRIPT_DIR))

import sap_bw_query_output_prettifier as prettifier   # fill_down_hierarchy
import sap_bw_compare_queries          as cq           # parse_excel, collect_elements, write_excel_report
import sap_bw_analyze_key_figures      as kf           # load_and_parse, extract_*, compare_*, write_excel_report


# ── Helpers ──────────────────────────────────────────────────────────────────

def _ask_file(prompt):
    """Prompt until the user supplies a valid .xlsx path."""
    while True:
        raw = input(prompt).strip().strip('"').strip("'")
        p   = Path(raw)
        if p.is_file() and p.suffix.lower() == '.xlsx':
            return str(p)
        if not p.exists():
            print(f'  Not found: {raw!r}')
        else:
            print(f'  Not an .xlsx file: {raw!r}')


def _copy_sheet(src_ws, dst_wb, title):
    """Copy a worksheet (values + basic styles + layout) into dst_wb."""
    dst_ws = dst_wb.create_sheet(title[:31])   # Excel sheet name limit

    for row in src_ws.iter_rows():
        for cell in row:
            new = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new.font          = copy(cell.font)
                new.fill          = copy(cell.fill)
                new.alignment     = copy(cell.alignment)
                new.border        = copy(cell.border)
                new.number_format = cell.number_format

    for col, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col].width = dim.width
    for row_idx, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[row_idx].height = dim.height
    for merge in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merge))

    dst_ws.freeze_panes = src_ws.freeze_panes
    if src_ws.auto_filter.ref:
        dst_ws.auto_filter.ref = src_ws.auto_filter.ref

    return dst_ws


def _add_data_sheet(src_wb, dst_wb, title):
    """Append a filled-hierarchy data sheet (values + column widths only)."""
    src_ws = src_wb.active
    dst_ws = dst_wb.create_sheet(title[:31])
    for row in src_ws.iter_rows(values_only=True):
        dst_ws.append(list(row))
    for col, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst_ws.column_dimensions[col].width = dim.width
    dst_ws.freeze_panes = src_ws.freeze_panes
    return dst_ws


# ── Main pipeline ─────────────────────────────────────────────────────────────

def run(file1, file2, out_path):
    name_a = Path(file1).stem
    name_b = Path(file2).stem

    # ── Step 1: Fill-down hierarchy (for display sheets) ─────────────────────
    print('\n[1/3] Filling hierarchy columns ...')
    wb_a = openpyxl.load_workbook(file1, data_only=True)
    prettifier.fill_down_hierarchy(wb_a.active)

    wb_b = openpyxl.load_workbook(file2, data_only=True)
    prettifier.fill_down_hierarchy(wb_b.active)
    print('      Done.')

    # ── Step 2: Element comparison ────────────────────────────────────────────
    print('[2/3] Running element comparison ...')
    nodes1 = cq.parse_excel(file1)
    nodes2 = cq.parse_excel(file2)
    root1  = cq.find_root(nodes1)
    root2  = cq.find_root(nodes2)
    set1   = cq.collect_elements(nodes1, root1)
    set2   = cq.collect_elements(nodes2, root2)

    ids1, ids2 = set(set1.keys()), set(set2.keys())
    print(f'      Elements — A: {len(ids1)}, B: {len(ids2)}, '
          f'common: {len(ids1 & ids2)}, only A: {len(ids1 - ids2)}, only B: {len(ids2 - ids1)}')

    # ── Step 3: Key figure analysis ───────────────────────────────────────────
    print('[3/3] Running key figure analysis ...')
    kf_nodes1 = kf.load_and_parse(file1)
    kf_nodes2 = kf.load_and_parse(file2)
    kf_root1  = kf.find_root(kf_nodes1)
    kf_root2  = kf.find_root(kf_nodes2)

    subs1 = kf.find_rows_columns_subtrees(kf_nodes1, kf_root1)
    subs2 = kf.find_rows_columns_subtrees(kf_nodes2, kf_root2)

    rkfs1 = kf.extract_rkfs(kf_nodes1, subs1)
    rkfs2 = kf.extract_rkfs(kf_nodes2, subs2)
    ckfs1 = kf.extract_ckfs(kf_nodes1, subs1)
    ckfs2 = kf.extract_ckfs(kf_nodes2, subs2)
    fmls1 = kf.extract_formulas(kf_nodes1, subs1)
    fmls2 = kf.extract_formulas(kf_nodes2, subs2)

    rkf_equiv, rkf_only_a, rkf_only_b = kf.compare_rkfs(rkfs1, rkfs2)
    ckf_pairs, ckf_only_a, ckf_only_b = kf.compare_ckfs(ckfs1, ckfs2)
    fml_pairs, fml_only_a, fml_only_b = kf.compare_formulas(fmls1, fmls2)

    print(f'      RKFs — A: {len(rkfs1)}, B: {len(rkfs2)}, equivalent: {len(rkf_equiv)}')
    print(f'      CKFs — A: {len(ckfs1)}, B: {len(ckfs2)}, matched: {len(ckf_pairs)}')
    print(f'      Formulas — A: {len(fmls1)}, B: {len(fmls2)}, matched: {len(fml_pairs)}')

    # ── Build combined workbook ───────────────────────────────────────────────
    print('\nAssembling combined workbook ...')
    combined = openpyxl.Workbook()
    combined.remove(combined.active)    # drop default empty sheet

    # Sheets 1–2: Filled hierarchy data
    _add_data_sheet(wb_a, combined, f'Query A (Filled)')
    _add_data_sheet(wb_b, combined, f'Query B (Filled)')
    print('  + Query A (Filled), Query B (Filled)')

    # Sheets 3–4: Comparison (generate in temp file, copy sheets)
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        tmp_cmp = tmp.name
    cq.write_excel_report(file1, file2, set1, set2, tmp_cmp)
    wb_cmp = openpyxl.load_workbook(tmp_cmp)
    cmp_rename = {'Summary': 'Comparison Summary'}
    for sname in wb_cmp.sheetnames:
        title = cmp_rename.get(sname, sname)
        _copy_sheet(wb_cmp[sname], combined, title)
        print(f'  + {title}')
    wb_cmp.close()
    os.unlink(tmp_cmp)

    # Sheets 5–8: KF analysis (generate in temp file, copy sheets)
    kf_args = (file1, file2,
               rkfs1, rkfs2, rkf_equiv, rkf_only_a, rkf_only_b,
               ckfs1, ckfs2, ckf_pairs, ckf_only_a, ckf_only_b,
               fmls1, fmls2, fml_pairs, fml_only_a, fml_only_b)
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        tmp_kf = tmp.name
    kf.write_excel_report(*kf_args, out_path=tmp_kf)
    wb_kf = openpyxl.load_workbook(tmp_kf)
    kf_rename = {'Summary': 'KF Summary'}
    for sname in wb_kf.sheetnames:
        title = kf_rename.get(sname, sname)
        _copy_sheet(wb_kf[sname], combined, title)
        print(f'  + {title}')
    wb_kf.close()
    os.unlink(tmp_kf)

    combined.save(out_path)
    print(f'\nCombined report saved: {out_path}')


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == '__main__':
    print('=' * 60)
    print('  SAP BW Query Analysis — Master Script')
    print('=' * 60)
    print()
    file1 = _ask_file('File A (.xlsx): ')
    file2 = _ask_file('File B (.xlsx): ')

    out_path = Path(file1).parent / 'bw_query_analysis_report.xlsx'
    run(file1, file2, str(out_path))
