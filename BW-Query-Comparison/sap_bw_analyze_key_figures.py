"""
Analyzes Restricted Key Figures (RKF), Calculated Key Figures (CKF),
and standalone Formulas (Row_XXXXX) in the Rows/Columns section of two
SAP BW query definition Excel exports.

Steps:
1. Fill-down hierarchical columns A–N in-memory (same as create_replica.py).
2. Build Nodekey/Parent tree.
3. Extract RKFs, CKFs, and Formulas from the Rows/Columns subtree
   (both Rows and Columns sub-sections).
4. Compare:
   - RKFs by structural content (Key Figure + Characteristic restrictions).
   - CKFs by fuzzy description match, then formula text comparison.
   - Formulas by fuzzy description match, then formula text comparison.
5. Write text report (kf_analysis_report.txt) and Excel report
   (kf_analysis_report.xlsx).
"""

import re
import sys
import difflib
import openpyxl
from pathlib import Path

from sap_bw_compare_queries import has_nodekey_columns, parse_indented

# ── Constants ────────────────────────────────────────────────────────────────

HIER_COLS      = list(range(1, 15))   # openpyxl 1-based: columns A–N
DATA_START_ROW = 4                     # rows 1–3 are header / title rows

# Column indices (1-based in openpyxl → 0-based in iter_rows tuples)
COL_DESC    = 15   # O  – Description
COL_VALUE   = 17   # Q  – Technical value / formula expression
COL_NODEKEY = 19   # S  – Nodekey (unique node ID)
COL_PARENT  = 20   # T  – Parent Nodekey
COL_DEFN    = 21   # U  – Definition type


# ── Phase 1: Fill-down ───────────────────────────────────────────────────────

def fill_down_hierarchy(ws):
    """Fill hierarchical columns A–N downward in-place (in-memory).

    When a new value appears in column X, columns X+1…N are reset.
    """
    current = {}
    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        new_col = None
        for col in HIER_COLS:
            v = ws.cell(row=row_idx, column=col).value
            if v is not None and str(v).strip():
                new_col = col
                break
        if new_col is not None:
            for col in HIER_COLS:
                if col > new_col:
                    current[col] = None
            current[new_col] = ws.cell(row=row_idx, column=new_col).value
        for col in HIER_COLS:
            v = current.get(col)
            if v is not None:
                ws.cell(row=row_idx, column=col).value = v


# ── Phase 2: Build tree ──────────────────────────────────────────────────────

def parse_worksheet(ws):
    """Build a Nodekey-keyed node dict from a (filled-down) worksheet."""
    nodes = {}
    for row in ws.iter_rows(values_only=True):
        nk_raw  = row[COL_NODEKEY - 1]
        par_raw = row[COL_PARENT  - 1]
        def_raw = row[COL_DEFN    - 1]
        dsc_raw = row[COL_DESC    - 1]
        val_raw = row[COL_VALUE   - 1]

        if nk_raw is None:
            continue

        nk   = str(nk_raw).strip()
        par  = str(par_raw).strip() if par_raw is not None else ''
        defn = str(def_raw).strip() if def_raw is not None else ''
        desc = str(dsc_raw).strip() if dsc_raw is not None else ''
        val  = str(val_raw).strip() if val_raw is not None else ''

        if not nk or not defn:
            continue

        nodes[nk] = {
            'parent':      par,
            'definition':  defn,
            'description': desc,
            'value':       val,
            'children':    [],
        }

    for nk, node in nodes.items():
        p = node['parent']
        if p and p in nodes:
            nodes[p]['children'].append(nk)

    return nodes


def find_root(nodes):
    for nk, node in nodes.items():
        if not node['parent'] or node['parent'] not in nodes:
            return nk
    return None


def load_and_parse(filepath):
    """Load workbook, apply fill-down in-memory, return node tree."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    if not has_nodekey_columns(ws):
        return parse_indented(ws)   # needs the raw indentation, so no fill-down
    fill_down_hierarchy(ws)
    return parse_worksheet(ws)


# ── Phase 3: Navigate to Rows/Columns ────────────────────────────────────────

def find_rows_columns_subtrees(nodes, root_key):
    """Return nodekeys for 'Rows' and 'Columns' children under Rows/Columns."""
    subtree_roots = []
    rc_node = None
    for child_nk in nodes[root_key]['children']:
        if nodes[child_nk]['definition'] == 'Rows/Columns':
            rc_node = child_nk
            break
    if rc_node is None:
        return subtree_roots
    for child_nk in nodes[rc_node]['children']:
        if nodes[child_nk]['definition'] in ('Rows', 'Columns'):
            subtree_roots.append(child_nk)
    return subtree_roots


# ── Phase 3: Extract objects ─────────────────────────────────────────────────

def _best_text(node):
    """Return the longer of description / value (formula text can be in either)."""
    desc, val = node['description'], node['value']
    if desc and val and desc != val:
        return desc if len(desc) >= len(val) else val
    return desc or val


def collect_formula_text(nodes, nk):
    """Collect formula expression text from a CKF/Formula node's subtree.

    Stops descent into Restricted KF subtrees — they are structural
    building blocks, not formula expression text.
    """
    parts = []

    def walk(k):
        node = nodes[k]
        if node['definition'] == 'Restricted KF':
            return
        if node['definition'] == 'Formula':
            text = _best_text(node)
            if text:
                parts.append(text)
        for child_nk in node['children']:
            walk(child_nk)

    for child_nk in nodes[nk]['children']:
        walk(child_nk)

    return ' '.join(parts).strip()


def build_rkf_definition(nodes, nk):
    """Return (key_figures, restrictions) for a Restricted KF node."""
    key_figures  = []
    restrictions = []

    for child_nk in nodes[nk]['children']:
        child = nodes[child_nk]
        defn  = child['definition']

        if defn == 'Key Figure':
            kf = child['value'] or child['description']
            if kf:
                key_figures.append(kf)

        elif defn == 'Key Figures':
            # Container node — collect Key Figure children inside
            for gc_nk in child['children']:
                gc = nodes[gc_nk]
                if gc['definition'] == 'Key Figure':
                    kf = gc['value'] or gc['description']
                    if kf:
                        key_figures.append(kf)

        elif defn == 'Characteristic':
            char_name = child['value'] or child['description']
            char_desc = child['description']
            ranges, variables = [], []
            for gc_nk in child['children']:
                gc = nodes[gc_nk]
                if gc['definition'] == 'Range':
                    v = gc['value'] or gc['description']
                    if v:
                        ranges.append(v)
                elif gc['definition'] == 'Variable':
                    v = gc['value'] or gc['description']
                    if v:
                        variables.append(v)
            restrictions.append({
                'char_name': char_name,
                'char_desc': char_desc,
                'ranges':    ranges,
                'variables': variables,
            })

    return key_figures, restrictions


def rkf_canonical(key_figures, restrictions):
    """Frozenset canonical form for structural RKF comparison."""
    kf_set = frozenset(key_figures)
    restr_set = frozenset(
        (r['char_name'], frozenset(r['ranges'] + r['variables']))
        for r in restrictions
    )
    return (kf_set, restr_set)


def extract_rkfs(nodes, subtree_roots):
    """Return {name: rkf_dict} for all Restricted KFs in the subtree."""
    rkfs = {}
    seen = set()

    def walk(nk):
        node = nodes[nk]
        if node['definition'] == 'Restricted KF' and node['value']:
            name = node['value']
            if name not in seen:
                seen.add(name)
                kfs, restrictions = build_rkf_definition(nodes, nk)
                rkfs[name] = {
                    'name':         name,
                    'description':  node['description'],
                    'key_figures':  kfs,
                    'restrictions': restrictions,
                    'canonical':    rkf_canonical(kfs, restrictions),
                }
        for child_nk in node['children']:
            walk(child_nk)

    for root in subtree_roots:
        walk(root)
    return rkfs


def extract_ckfs(nodes, subtree_roots):
    """Return {name: ckf_dict} for all Calculated KFs (deduped by name)."""
    ckfs = {}
    seen = set()

    def walk(nk):
        node = nodes[nk]
        if node['definition'] == 'Calculated KF' and node['value']:
            name = node['value']
            if name not in seen:
                seen.add(name)
                ckfs[name] = {
                    'name':         name,
                    'description':  node['description'],
                    'formula_text': collect_formula_text(nodes, nk),
                }
        for child_nk in node['children']:
            walk(child_nk)

    for root in subtree_roots:
        walk(root)
    return ckfs


def extract_formulas(nodes, subtree_roots):
    """Return {row_id: formula_dict} for standalone Formulas (Row_XXXXX)
    whose immediate parent has definition='Structure'.
    """
    formulas = {}

    def walk(nk, parent_defn=None):
        node = nodes[nk]
        if (node['definition'] == 'Formula'
                and node['value'].startswith('Row_')
                and parent_defn == 'Structure'):
            row_id = node['value']
            if row_id not in formulas:
                formulas[row_id] = {
                    'row_id':       row_id,
                    'description':  node['description'],
                    'formula_text': collect_formula_text(nodes, nk),
                }
        for child_nk in node['children']:
            walk(child_nk, node['definition'])

    for root in subtree_roots:
        walk(root)
    return formulas


# ── Phase 4: Comparison ──────────────────────────────────────────────────────

def normalize(text):
    """Normalize whitespace for formula text comparison."""
    return re.sub(r'\s+', ' ', text or '').strip()


def fuzzy_match_by_desc(items_a, items_b, desc_fn, threshold=0.80):
    """Match items from A to B by fuzzy description similarity.

    Returns (matched [(key_a, key_b, ratio)], only_a [keys], only_b [keys]).
    """
    descs_a = {k: desc_fn(v).lower() for k, v in items_a.items()}
    descs_b = {k: desc_fn(v).lower() for k, v in items_b.items()}

    matched, used_b, only_a = [], set(), []

    for k_a, desc_a in descs_a.items():
        if not desc_a:
            only_a.append(k_a)
            continue
        best_ratio, best_k_b = 0.0, None
        for k_b, desc_b in descs_b.items():
            if k_b in used_b or not desc_b:
                continue
            r = difflib.SequenceMatcher(None, desc_a, desc_b).ratio()
            if r > best_ratio:
                best_ratio, best_k_b = r, k_b
        if best_ratio >= threshold and best_k_b is not None:
            matched.append((k_a, best_k_b, best_ratio))
            used_b.add(best_k_b)
        else:
            only_a.append(k_a)

    only_b = [k for k in descs_b if k not in used_b]
    return matched, only_a, only_b


def compare_rkfs(rkfs_a, rkfs_b):
    """Match RKFs by canonical structural form.

    Returns (equivalent [(na, nb)], only_a [names], only_b [names]).
    """
    canon_a = {}
    for name, rkf in rkfs_a.items():
        canon_a.setdefault(rkf['canonical'], []).append(name)
    canon_b = {}
    for name, rkf in rkfs_b.items():
        canon_b.setdefault(rkf['canonical'], []).append(name)

    all_canons = set(canon_a) | set(canon_b)
    equivalent, only_a, only_b = [], [], []

    for canon in all_canons:
        names_a = canon_a.get(canon, [])
        names_b = canon_b.get(canon, [])
        for i in range(max(len(names_a), len(names_b))):
            na = names_a[i] if i < len(names_a) else None
            nb = names_b[i] if i < len(names_b) else None
            if na and nb:
                equivalent.append((na, nb))
            elif na:
                only_a.append(na)
            else:
                only_b.append(nb)

    return equivalent, only_a, only_b


def compare_ckfs(ckfs_a, ckfs_b):
    """Fuzzy-match CKFs by description, compare formula text."""
    matched, only_a, only_b = fuzzy_match_by_desc(
        ckfs_a, ckfs_b, lambda x: x['description'])
    pairs = []
    for k_a, k_b, ratio in matched:
        same = normalize(ckfs_a[k_a]['formula_text']) == normalize(ckfs_b[k_b]['formula_text'])
        pairs.append({
            'name_a': k_a, 'name_b': k_b, 'ratio': ratio,
            'same_formula': same,
            'ckf_a': ckfs_a[k_a], 'ckf_b': ckfs_b[k_b],
        })
    return pairs, only_a, only_b


def compare_formulas(fmls_a, fmls_b):
    """Fuzzy-match formulas by description, compare formula text."""
    matched, only_a, only_b = fuzzy_match_by_desc(
        fmls_a, fmls_b, lambda x: x['description'])
    pairs = []
    for k_a, k_b, ratio in matched:
        same = normalize(fmls_a[k_a]['formula_text']) == normalize(fmls_b[k_b]['formula_text'])
        pairs.append({
            'row_a': k_a, 'row_b': k_b, 'ratio': ratio,
            'same_formula': same,
            'fml_a': fmls_a[k_a], 'fml_b': fmls_b[k_b],
        })
    return pairs, only_a, only_b


# ── Phase 5: Text report ──────────────────────────────────────────────────────

def _rkf_lines(rkf):
    lines = []
    if rkf['key_figures']:
        lines.append('    Key Figure(s): ' + ', '.join(rkf['key_figures']))
    for r in rkf['restrictions']:
        char = r['char_name']
        if r['char_desc'] and r['char_desc'] != r['char_name']:
            char += f" ({r['char_desc']})"
        vals = ', '.join(r['ranges'] + r['variables'])
        lines.append(f"    {char}: {vals}" if vals else f"    {char}")
    return lines


def write_text_report(stream, file1, file2,
                      rkfs_a, rkfs_b, rkf_equiv, rkf_only_a, rkf_only_b,
                      ckfs_a, ckfs_b, ckf_pairs, ckf_only_a, ckf_only_b,
                      fmls_a, fmls_b, fml_pairs, fml_only_a, fml_only_b):
    W = 80

    def sep(c='='):
        stream.write(c * W + '\n')

    def ln(*a):
        stream.write(' '.join(str(x) for x in a) + '\n')

    sep()
    ln('KEY FIGURE ANALYSIS REPORT')
    sep()
    ln(f'  File A : {Path(file1).name}')
    ln(f'  File B : {Path(file2).name}')
    ln()

    # ── Restricted Key Figures ────────────────────────────────────────────────
    sep()
    ln('RESTRICTED KEY FIGURES')
    sep()
    ln(f'  In A: {len(rkfs_a)}  |  In B: {len(rkfs_b)}  |  '
       f'Equivalent: {len(rkf_equiv)}  |  Only A: {len(rkf_only_a)}  |  '
       f'Only B: {len(rkf_only_b)}')
    ln()

    if rkf_equiv:
        ln('  ── STRUCTURALLY EQUIVALENT')
        for na, nb in sorted(rkf_equiv, key=lambda x: rkfs_a[x[0]]['description']):
            rkf = rkfs_a[na]
            ln(f'  A: {na}  ({rkf["description"]})')
            ln(f'  B: {nb}  ({rkfs_b[nb]["description"]})')
            for line in _rkf_lines(rkf):
                ln(line)
            ln()

    if rkf_only_a:
        sep('-')
        ln('  ONLY IN FILE A')
        sep('-')
        for name in sorted(rkf_only_a, key=lambda x: rkfs_a[x]['description']):
            rkf = rkfs_a[name]
            ln(f'  {name}  ({rkf["description"]})')
            for line in _rkf_lines(rkf):
                ln(line)
            ln()

    if rkf_only_b:
        sep('-')
        ln('  ONLY IN FILE B')
        sep('-')
        for name in sorted(rkf_only_b, key=lambda x: rkfs_b[x]['description']):
            rkf = rkfs_b[name]
            ln(f'  {name}  ({rkf["description"]})')
            for line in _rkf_lines(rkf):
                ln(line)
            ln()

    # ── Calculated Key Figures ────────────────────────────────────────────────
    sep()
    ln('CALCULATED KEY FIGURES')
    sep()
    same_ckf = sum(1 for p in ckf_pairs if p['same_formula'])
    ln(f'  In A: {len(ckfs_a)}  |  In B: {len(ckfs_b)}  |  '
       f'Matched: {len(ckf_pairs)} (same formula: {same_ckf}, '
       f'different: {len(ckf_pairs) - same_ckf})  |  '
       f'Only A: {len(ckf_only_a)}  |  Only B: {len(ckf_only_b)}')
    ln()

    if ckf_pairs:
        ln('  ── MATCHED (by description name)')
        for p in sorted(ckf_pairs, key=lambda x: x['ckf_a']['description']):
            sym = '=' if p['same_formula'] else '≠'
            ln(f'  {sym} A: {p["name_a"]}  ({p["ckf_a"]["description"]})')
            ln(f'    B: {p["name_b"]}  ({p["ckf_b"]["description"]})')
            if not p['same_formula']:
                ln(f'    Formula A: {p["ckf_a"]["formula_text"][:150]}')
                ln(f'    Formula B: {p["ckf_b"]["formula_text"][:150]}')
            ln()

    if ckf_only_a:
        sep('-')
        ln('  ONLY IN FILE A')
        sep('-')
        for name in sorted(ckf_only_a, key=lambda x: ckfs_a[x]['description']):
            ckf = ckfs_a[name]
            ln(f'  {name}  ({ckf["description"]})')
            if ckf['formula_text']:
                ln(f'    {ckf["formula_text"][:150]}')
            ln()

    if ckf_only_b:
        sep('-')
        ln('  ONLY IN FILE B')
        sep('-')
        for name in sorted(ckf_only_b, key=lambda x: ckfs_b[x]['description']):
            ckf = ckfs_b[name]
            ln(f'  {name}  ({ckf["description"]})')
            if ckf['formula_text']:
                ln(f'    {ckf["formula_text"][:150]}')
            ln()

    # ── Standalone Formulas ───────────────────────────────────────────────────
    sep()
    ln('STANDALONE FORMULAS  (Row_XXXXX)')
    sep()
    same_fml = sum(1 for p in fml_pairs if p['same_formula'])
    ln(f'  In A: {len(fmls_a)}  |  In B: {len(fmls_b)}  |  '
       f'Matched: {len(fml_pairs)} (same: {same_fml}, '
       f'different: {len(fml_pairs) - same_fml})  |  '
       f'Only A: {len(fml_only_a)}  |  Only B: {len(fml_only_b)}')
    ln()

    if fml_pairs:
        ln('  ── MATCHED (by description name)')
        for p in sorted(fml_pairs, key=lambda x: x['fml_a']['description']):
            sym = '=' if p['same_formula'] else '≠'
            ln(f'  {sym} A: {p["row_a"]}  ({p["fml_a"]["description"]})')
            ln(f'    B: {p["row_b"]}  ({p["fml_b"]["description"]})')
            if not p['same_formula']:
                ln(f'    Formula A: {p["fml_a"]["formula_text"][:150]}')
                ln(f'    Formula B: {p["fml_b"]["formula_text"][:150]}')
            ln()

    if fml_only_a:
        sep('-')
        ln('  ONLY IN FILE A')
        sep('-')
        for rid in sorted(fml_only_a, key=lambda x: fmls_a[x]['description']):
            f = fmls_a[rid]
            ln(f'  {rid}  ({f["description"]})')
            if f['formula_text']:
                ln(f'    {f["formula_text"][:150]}')
            ln()

    if fml_only_b:
        sep('-')
        ln('  ONLY IN FILE B')
        sep('-')
        for rid in sorted(fml_only_b, key=lambda x: fmls_b[x]['description']):
            f = fmls_b[rid]
            ln(f'  {rid}  ({f["description"]})')
            if f['formula_text']:
                ln(f'    {f["formula_text"][:150]}')
            ln()

    sep()


# ── Phase 5: Excel report ─────────────────────────────────────────────────────

def write_excel_report(file1, file2,
                       rkfs_a, rkfs_b, rkf_equiv, rkf_only_a, rkf_only_b,
                       ckfs_a, ckfs_b, ckf_pairs, ckf_only_a, ckf_only_b,
                       fmls_a, fmls_b, fml_pairs, fml_only_a, fml_only_b,
                       out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    C_HEADER = '1F4E79'
    C_EQUIV  = 'FFFFFF'   # white  – structurally equivalent / same formula
    C_DIFF   = 'FFF2CC'   # yellow – matched but different formula
    C_ONLY_A = 'FCE4D6'   # orange – only in A
    C_ONLY_B = 'E2EFDA'   # green  – only in B

    thin   = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    name_a = Path(file1).stem[:28]
    name_b = Path(file2).stem[:28]

    wb = Workbook()

    # ── helpers ───────────────────────────────────────────────────────────────

    def header_row(ws, headers):
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.font      = Font(bold=True, color='FFFFFF', size=10)
            c.fill      = PatternFill('solid', fgColor=C_HEADER)
            c.alignment = Alignment(horizontal='center', vertical='center',
                                    wrap_text=True)
            c.border    = border
        ws.row_dimensions[1].height = 36

    def style_data_row(ws, row_num, n_cols, fill_color):
        fill = PatternFill('solid', fgColor=fill_color)
        for col in range(1, n_cols + 1):
            c = ws.cell(row=row_num, column=col)
            c.fill      = fill
            c.border    = border
            c.alignment = Alignment(horizontal='left', vertical='top',
                                    wrap_text=True)
        ws.row_dimensions[row_num].height = 60

    def fmt_restrictions(rkf):
        parts = []
        if rkf['key_figures']:
            parts.append('KF: ' + ', '.join(rkf['key_figures']))
        for r in rkf['restrictions']:
            char = r['char_name']
            if r['char_desc'] and r['char_desc'] != r['char_name']:
                char += f" ({r['char_desc']})"
            vals = ', '.join(r['ranges'] + r['variables'])
            parts.append(f"{char}: {vals}" if vals else char)
        return '\n'.join(parts)

    # ── Restricted KFs sheet ──────────────────────────────────────────────────
    ws_rkf = wb.active
    ws_rkf.title = 'Restricted KFs'
    header_row(ws_rkf, [
        f'Name A\n({name_a[:22]})', 'Desc A',
        f'Name B\n({name_b[:22]})', 'Desc B',
        'Key Figure(s)', 'Restrictions', 'Status',
    ])
    rn = 2

    for na, nb in sorted(rkf_equiv, key=lambda x: rkfs_a[x[0]]['description']):
        rkf = rkfs_a[na]
        ws_rkf.append([na, rkf['description'], nb, rkfs_b[nb]['description'],
                       ', '.join(rkf['key_figures']), fmt_restrictions(rkf),
                       'Equivalent'])
        style_data_row(ws_rkf, rn, 7, C_EQUIV)
        rn += 1

    for name in sorted(rkf_only_a, key=lambda x: rkfs_a[x]['description']):
        rkf = rkfs_a[name]
        ws_rkf.append([name, rkf['description'], '', '',
                       ', '.join(rkf['key_figures']), fmt_restrictions(rkf),
                       'Only in A'])
        style_data_row(ws_rkf, rn, 7, C_ONLY_A)
        rn += 1

    for name in sorted(rkf_only_b, key=lambda x: rkfs_b[x]['description']):
        rkf = rkfs_b[name]
        ws_rkf.append(['', '', name, rkf['description'],
                       ', '.join(rkf['key_figures']), fmt_restrictions(rkf),
                       'Only in B'])
        style_data_row(ws_rkf, rn, 7, C_ONLY_B)
        rn += 1

    for i, w in enumerate([38, 38, 38, 38, 22, 65, 12], 1):
        ws_rkf.column_dimensions[get_column_letter(i)].width = w
    ws_rkf.freeze_panes = 'A2'
    ws_rkf.auto_filter.ref = f'A1:G{rn - 1}'

    # ── Calculated KFs sheet ──────────────────────────────────────────────────
    ws_ckf = wb.create_sheet('Calculated KFs')
    header_row(ws_ckf, [
        f'Name A\n({name_a[:22]})', 'Desc A', 'Formula A',
        f'Name B\n({name_b[:22]})', 'Desc B', 'Formula B',
        'Formula\nMatch?', 'Status',
    ])
    rn = 2

    for p in sorted(ckf_pairs, key=lambda x: x['ckf_a']['description']):
        fill   = C_EQUIV if p['same_formula'] else C_DIFF
        status = 'Equivalent' if p['same_formula'] else 'Different Formula'
        ws_ckf.append([
            p['name_a'], p['ckf_a']['description'], p['ckf_a']['formula_text'],
            p['name_b'], p['ckf_b']['description'], p['ckf_b']['formula_text'],
            '✓' if p['same_formula'] else '✗', status,
        ])
        style_data_row(ws_ckf, rn, 8, fill)
        rn += 1

    for name in sorted(ckf_only_a, key=lambda x: ckfs_a[x]['description']):
        ckf = ckfs_a[name]
        ws_ckf.append([name, ckf['description'], ckf['formula_text'],
                       '', '', '', '', 'Only in A'])
        style_data_row(ws_ckf, rn, 8, C_ONLY_A)
        rn += 1

    for name in sorted(ckf_only_b, key=lambda x: ckfs_b[x]['description']):
        ckf = ckfs_b[name]
        ws_ckf.append(['', '', '', name, ckf['description'], ckf['formula_text'],
                       '', 'Only in B'])
        style_data_row(ws_ckf, rn, 8, C_ONLY_B)
        rn += 1

    for i, w in enumerate([38, 38, 70, 38, 38, 70, 10, 18], 1):
        ws_ckf.column_dimensions[get_column_letter(i)].width = w
    ws_ckf.freeze_panes = 'A2'
    ws_ckf.auto_filter.ref = f'A1:H{rn - 1}'

    # ── Formulas sheet ────────────────────────────────────────────────────────
    ws_fml = wb.create_sheet('Formulas')
    header_row(ws_fml, [
        f'Row ID A\n({name_a[:22]})', 'Desc A', 'Formula A',
        f'Row ID B\n({name_b[:22]})', 'Desc B', 'Formula B',
        'Formula\nMatch?', 'Status',
    ])
    rn = 2

    for p in sorted(fml_pairs, key=lambda x: x['fml_a']['description']):
        fill   = C_EQUIV if p['same_formula'] else C_DIFF
        status = 'Equivalent' if p['same_formula'] else 'Different Formula'
        ws_fml.append([
            p['row_a'], p['fml_a']['description'], p['fml_a']['formula_text'],
            p['row_b'], p['fml_b']['description'], p['fml_b']['formula_text'],
            '✓' if p['same_formula'] else '✗', status,
        ])
        style_data_row(ws_fml, rn, 8, fill)
        rn += 1

    for rid in sorted(fml_only_a, key=lambda x: fmls_a[x]['description']):
        f = fmls_a[rid]
        ws_fml.append([rid, f['description'], f['formula_text'],
                       '', '', '', '', 'Only in A'])
        style_data_row(ws_fml, rn, 8, C_ONLY_A)
        rn += 1

    for rid in sorted(fml_only_b, key=lambda x: fmls_b[x]['description']):
        f = fmls_b[rid]
        ws_fml.append(['', '', '', rid, f['description'], f['formula_text'],
                       '', 'Only in B'])
        style_data_row(ws_fml, rn, 8, C_ONLY_B)
        rn += 1

    for i, w in enumerate([16, 42, 90, 16, 42, 90, 10, 18], 1):
        ws_fml.column_dimensions[get_column_letter(i)].width = w
    ws_fml.freeze_panes = 'A2'
    ws_fml.auto_filter.ref = f'A1:H{rn - 1}'

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws_sum = wb.create_sheet('Summary')
    ws_sum.column_dimensions['A'].width = 44
    ws_sum.column_dimensions['B'].width = 16

    same_ckf = sum(1 for p in ckf_pairs if p['same_formula'])
    same_fml = sum(1 for p in fml_pairs if p['same_formula'])

    summary_rows = [
        ('File A', Path(file1).name),
        ('File B', Path(file2).name),
        ('', ''),
        ('RESTRICTED KEY FIGURES', ''),
        ('  In File A',                    len(rkfs_a)),
        ('  In File B',                    len(rkfs_b)),
        ('  Structurally Equivalent',      len(rkf_equiv)),
        ('  Only in A',                    len(rkf_only_a)),
        ('  Only in B',                    len(rkf_only_b)),
        ('', ''),
        ('CALCULATED KEY FIGURES', ''),
        ('  In File A',                    len(ckfs_a)),
        ('  In File B',                    len(ckfs_b)),
        ('  Matched — same formula',       same_ckf),
        ('  Matched — different formula',  len(ckf_pairs) - same_ckf),
        ('  Only in A',                    len(ckf_only_a)),
        ('  Only in B',                    len(ckf_only_b)),
        ('', ''),
        ('STANDALONE FORMULAS', ''),
        ('  In File A',                    len(fmls_a)),
        ('  In File B',                    len(fmls_b)),
        ('  Matched — same formula',       same_fml),
        ('  Matched — different formula',  len(fml_pairs) - same_fml),
        ('  Only in A',                    len(fml_only_a)),
        ('  Only in B',                    len(fml_only_b)),
    ]

    section_fill = PatternFill('solid', fgColor='D6E4F0')
    for r, (label, val) in enumerate(summary_rows, 1):
        is_section = bool(label and not label.startswith('  '))
        ca = ws_sum.cell(row=r, column=1, value=label)
        cb = ws_sum.cell(row=r, column=2, value=val if val != '' else None)
        ca.font = Font(bold=is_section)
        if is_section and label:
            ca.fill = section_fill
            cb.fill = section_fill

    wb.save(out_path)


# ── Entry point ───────────────────────────────────────────────────────────────

def analyze(file1, file2):
    print(f'Parsing  : {Path(file1).name}')
    nodes1 = load_and_parse(file1)
    root1  = find_root(nodes1)

    print(f'Parsing  : {Path(file2).name}')
    nodes2 = load_and_parse(file2)
    root2  = find_root(nodes2)

    if root1 is None or root2 is None:
        print('ERROR: Could not find root node.')
        return

    subs1 = find_rows_columns_subtrees(nodes1, root1)
    subs2 = find_rows_columns_subtrees(nodes2, root2)
    print(f'  Rows/Columns subtrees found: A={len(subs1)}, B={len(subs2)}')

    rkfs1 = extract_rkfs(nodes1, subs1)
    rkfs2 = extract_rkfs(nodes2, subs2)
    ckfs1 = extract_ckfs(nodes1, subs1)
    ckfs2 = extract_ckfs(nodes2, subs2)
    fmls1 = extract_formulas(nodes1, subs1)
    fmls2 = extract_formulas(nodes2, subs2)

    print(f'  Extracted — RKFs: A={len(rkfs1)}, B={len(rkfs2)}  |  '
          f'CKFs: A={len(ckfs1)}, B={len(ckfs2)}  |  '
          f'Formulas: A={len(fmls1)}, B={len(fmls2)}')

    rkf_equiv, rkf_only_a, rkf_only_b = compare_rkfs(rkfs1, rkfs2)
    ckf_pairs, ckf_only_a, ckf_only_b = compare_ckfs(ckfs1, ckfs2)
    fml_pairs, fml_only_a, fml_only_b = compare_formulas(fmls1, fmls2)

    txt_path  = Path(file1).parent / 'kf_analysis_report.txt'
    xlsx_path = Path(file1).parent / 'kf_analysis_report.xlsx'

    common = (file1, file2,
              rkfs1, rkfs2, rkf_equiv, rkf_only_a, rkf_only_b,
              ckfs1, ckfs2, ckf_pairs, ckf_only_a, ckf_only_b,
              fmls1, fmls2, fml_pairs, fml_only_a, fml_only_b)

    write_text_report(sys.stdout, *common)

    with open(txt_path, 'w', encoding='utf-8') as fh:
        write_text_report(fh, *common)

    write_excel_report(*common, out_path=xlsx_path)

    print(f'\nText report : {txt_path}')
    print(f'Excel report: {xlsx_path}')


if __name__ == '__main__':
    base  = Path(__file__).parent.parent / 'Query Comparison'
    files = sorted(f for f in base.glob('G_*.xlsx') if '_replica' not in f.name)
    if len(files) != 2:
        print(f'Expected exactly 2 .xlsx files, found: {files}')
    else:
        analyze(str(files[0]), str(files[1]))
